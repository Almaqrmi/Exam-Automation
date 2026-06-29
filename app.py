from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.utils import secure_filename
import os
import json
from datetime import datetime
from io import BytesIO
import zipfile
import random

# استيراد الإعدادات والنماذج
from config import Config
from models import db, User, Exam, Question

# استيراد مكتبات معالجة PDF والذكاء الاصطناعي
import PyPDF2
import google.generativeai as genai

# استيراد مكتبات التصدير
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxtpl import DocxTemplate
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT

# إنشاء التطبيق
app = Flask(__name__)
app.config.from_object(Config)
Config.init_app(app)

# إضافة فلتر from_json لـ Jinja2
@app.template_filter('from_json')
def from_json_filter(value):
    """تحويل JSON string إلى Python object"""
    if value:
        try:
            return json.loads(value)
        except:
            return []
    return []

# تهيئة قاعدة البيانات
db.init_app(app)

# تهيئة نظام تسجيل الدخول
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message = 'يرجى تسجيل الدخول للوصول إلى هذه الصفحة'

# تكوين Gemini API
if app.config['GEMINI_API_KEY']:
    genai.configure(api_key=app.config['GEMINI_API_KEY'])


@login_manager.user_loader
def load_user(user_id):
    """تحميل المستخدم من قاعدة البيانات"""
    return User.query.get(int(user_id))


def allowed_file(filename):
    """التحقق من امتداد الملف المسموح"""
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']


def extract_text_from_pdf(pdf_path):
    """استخراج النص من ملف PDF"""
    try:
        with open(pdf_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
    except Exception as e:
        print(f"خطأ في قراءة PDF: {str(e)}")
        return None


def generate_questions_with_ai(text, num_questions, question_type, difficulty):
    """Generate questions using Gemini AI"""
    try:
        # استخدم النموذج الذي يعمل
        model = genai.GenerativeModel('gemini-2.5-flash')
        
        # Build the prompt based on question type
        type_mapping = {
            'mixed': 'Mixed (Multiple Choice, True/False, Fill in the Blank)',
            'mcq': 'Multiple Choice only',
            'tf': 'True/False only',
            'fill': 'Fill in the Blank only',
            'direct': 'Direct questions only'
        }
        
        difficulty_mapping = {
            'easy': 'Easy',
            'medium': 'Medium',
            'hard': 'Hard'
        }
        
        prompt = f"""
You are an educational assistant specialized in creating exams.

Educational Content:
{text[:4000]}

Requirements:
- Number of questions: {num_questions}
- Question type: {type_mapping.get(question_type, 'Mixed')}
- Difficulty level: {difficulty_mapping.get(difficulty, 'Medium')}

Create the questions and return the result in JSON format only without any additional text, in this format:
{{
  "questions": [
    {{
      "question": "Question text",
      "type": "mcq or tf or fill or direct",
      "answer": "Correct answer",
      "options": ["Option1", "Option2", "Option3", "Option4"]
    }}
  ]
}}

Notes:
- For mcq questions: must have 4 options and answer is one of them
- For tf questions: answer must be "True" or "False" and options should be null
- Make sure questions are clear and related to the content
- All questions must be in ENGLISH
"""
        
        # Generation settings
        generation_config = {
            'temperature': 0.7,
            'top_p': 0.95,
            'top_k': 40,
            'max_output_tokens': 8192,
        }
        
        response = model.generate_content(
            prompt,
            generation_config=generation_config
        )
        
        result_text = response.text.strip()
        
        # Clean text from any markdown
        if result_text.startswith('```json'):
            result_text = result_text[7:]
        if result_text.startswith('```'):
            result_text = result_text[3:]
        if result_text.endswith('```'):
            result_text = result_text[:-3]
        
        result_text = result_text.strip()
        
        # Convert result to JSON
        questions_data = json.loads(result_text)
        return questions_data
        
    except json.JSONDecodeError as e:
        print(f"JSON parsing error: {str(e)}")
        print(f"Received text: {result_text[:500] if 'result_text' in locals() else 'None'}")
        return None
    except Exception as e:
        print(f"Error generating questions: {str(e)}")
        print(f"Error type: {type(e).__name__}")
        return None


# ================== المسارات (Routes) ==================

@app.route('/')
def index():
    """الصفحة الرئيسية"""
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))


@app.route('/register', methods=['GET', 'POST'])
def register():
    """صفحة التسجيل"""
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        username = request.form.get('username')
        email = request.form.get('email')
        password = request.form.get('password')
        full_name = request.form.get('full_name')
        
        # التحقق من البيانات
        if not all([username, email, password]):
            flash('جميع الحقول مطلوبة', 'error')
            return redirect(url_for('register'))
        
        # التحقق من وجود المستخدم
        if User.query.filter_by(username=username).first():
            flash('اسم المستخدم موجود بالفعل', 'error')
            return redirect(url_for('register'))
        
        if User.query.filter_by(email=email).first():
            flash('البريد الإلكتروني مسجل بالفعل', 'error')
            return redirect(url_for('register'))
        
        # إنشاء مستخدم جديد
        user = User(username=username, email=email, full_name=full_name)
        user.set_password(password)
        
        db.session.add(user)
        db.session.commit()
        
        flash('تم إنشاء الحساب بنجاح! يمكنك الآن تسجيل الدخول', 'success')
        return redirect(url_for('login'))
    
    return render_template('register.html')


@app.route('/login', methods=['GET', 'POST'])
def login():
    """صفحة تسجيل الدخول"""
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        
        user = User.query.filter_by(username=username).first()
        
        if user and user.check_password(password):
            login_user(user)
            flash('مرحباً بك!', 'success')
            next_page = request.args.get('next')
            return redirect(next_page or url_for('dashboard'))
        else:
            flash('اسم المستخدم أو كلمة المرور غير صحيحة', 'error')
    
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    """تسجيل الخروج"""
    logout_user()
    flash('تم تسجيل الخروج بنجاح', 'success')
    return redirect(url_for('login'))


@app.route('/dashboard')
@login_required
def dashboard():
    """لوحة التحكم الرئيسية"""
    # جلب اختبارات المستخدم
    exams = Exam.query.filter_by(user_id=current_user.id).order_by(Exam.created_at.desc()).all()
    return render_template('dashboard.html', exams=exams)


@app.route('/generate', methods=['POST'])
@login_required
def generate_exam():
    """توليد اختبار جديد من ملف أو عدة ملفات"""
    try:
        # التحقق من وجود ملفات
        if 'pdf_files' not in request.files:
            return jsonify({'error': 'No files uploaded'}), 400
        
        files = request.files.getlist('pdf_files')
        
        if not files or len(files) == 0:
            return jsonify({'error': 'No files selected'}), 400
        
        # التحقق من عدد الملفات
        if len(files) > 5:
            return jsonify({'error': 'Maximum 5 files allowed'}), 400
        
        # قائمة لحفظ أسماء الملفات والنصوص
        saved_filenames = []
        all_texts = []
        
        # معالجة كل ملف
        for file in files:
            if file.filename == '':
                continue
                
            if not allowed_file(file.filename):
                return jsonify({'error': f'Invalid file type: {file.filename}. Only PDF allowed'}), 400
            
            # حفظ الملف
            filename = secure_filename(file.filename)
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{timestamp}_{filename}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(filepath)
            
            saved_filenames.append(filename)
            
            # استخراج النص من PDF
            text = extract_text_from_pdf(filepath)
            
            if not text:
                # حذف الملفات المحفوظة في حالة الفشل
                for fn in saved_filenames:
                    fp = os.path.join(app.config['UPLOAD_FOLDER'], fn)
                    if os.path.exists(fp):
                        os.remove(fp)
                return jsonify({'error': f'Cannot read content from: {file.filename}'}), 400
            
            all_texts.append(text)
        
        # دمج جميع النصوص
        combined_text = "\n\n--- NEW DOCUMENT ---\n\n".join(all_texts)
        
        # الحصول على معاملات التوليد
        num_questions = int(request.form.get('num_questions', 10))
        question_type = request.form.get('question_type', 'mixed')
        difficulty = request.form.get('difficulty', 'medium')
        exam_title = request.form.get('exam_title', 'New Exam')
        
        # الإعدادات الإضافية الجديدة
        show_answers = request.form.get('show_answers') == 'on'
        num_versions = int(request.form.get('num_versions', 1))
        teacher_name = request.form.get('teacher_name', '')
        
        # توليد الأسئلة
        questions_data = generate_questions_with_ai(combined_text, num_questions, question_type, difficulty)
        
        if not questions_data or 'questions' not in questions_data:
            # حذف الملفات في حالة الفشل
            for filename in saved_filenames:
                filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                if os.path.exists(filepath):
                    os.remove(filepath)
            return jsonify({'error': 'Failed to generate questions'}), 500
        
        # حفظ الاختبار في قاعدة البيانات
        exam = Exam(
            title=exam_title,
            pdf_filename=saved_filenames[0] if len(saved_filenames) == 1 else None,  # للتوافق مع القديم
            pdf_files=json.dumps(saved_filenames),  # حفظ جميع الملفات
            num_questions=num_questions,
            question_type=question_type,
            difficulty=difficulty,
            show_answers=show_answers,
            num_versions=num_versions,
            teacher_name=teacher_name,
            user_id=current_user.id
        )
        db.session.add(exam)
        db.session.commit()
        
        # حفظ الأسئلة
        for q_data in questions_data['questions']:
            question = Question(
                question_text=q_data['question'],
                question_type=q_data.get('type', 'direct'),
                correct_answer=q_data.get('answer', ''),
                options=json.dumps(q_data.get('options', [])) if q_data.get('options') else None,
                exam_id=exam.id
            )
            db.session.add(question)
        
        db.session.commit()
        
        return jsonify({
            'success': True,
            'exam_id': exam.id,
            'files_count': len(saved_filenames),
            'redirect': url_for('view_exam', exam_id=exam.id)
        })
        
    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({'error': 'An error occurred while generating the exam'}), 500


@app.route('/exam/<int:exam_id>')
@login_required
def view_exam(exam_id):
    """عرض الاختبار والأسئلة"""
    exam = Exam.query.get_or_404(exam_id)
    
    # التأكد من أن المستخدم يملك هذا الاختبار
    if exam.user_id != current_user.id:
        flash('غير مصرح لك بعرض هذا الاختبار', 'error')
        return redirect(url_for('dashboard'))
    
    questions = Question.query.filter_by(exam_id=exam_id).all()
    
    # تحويل options من JSON إلى list
    for q in questions:
        if q.options:
            q.options = json.loads(q.options)
    
    return render_template('results.html', exam=exam, questions=questions)


@app.route('/exam/<int:exam_id>/delete', methods=['POST'])
@login_required
def delete_exam(exam_id):
    """حذف اختبار"""
    exam = Exam.query.get_or_404(exam_id)
    
    if exam.user_id != current_user.id:
        return jsonify({'error': 'غير مصرح'}), 403
    
    # حذف ملفات PDF
    if exam.pdf_files:
        # حذف الملفات المتعددة
        try:
            filenames = json.loads(exam.pdf_files)
            for filename in filenames:
                filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                if os.path.exists(filepath):
                    os.remove(filepath)
        except:
            pass
    elif exam.pdf_filename:
        # حذف الملف الواحد (للتوافق مع الاختبارات القديمة)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], exam.pdf_filename)
        if os.path.exists(filepath):
            os.remove(filepath)
    
    db.session.delete(exam)
    db.session.commit()
    
    flash('تم حذف الاختبار بنجاح', 'success')
    return redirect(url_for('dashboard'))


@app.route('/exam/<int:exam_id>/export/word')
@login_required
def export_word(exam_id):
    """Export exam to Word format using template"""
    exam = Exam.query.get_or_404(exam_id)
    
    if exam.user_id != current_user.id:
        flash('Not authorized to export this exam', 'error')
        return redirect(url_for('dashboard'))
    
    questions = Question.query.filter_by(exam_id=exam_id).all()
    
    # Prepare questions data
    questions_data = []
    for q in questions:
        q_data = {
            'question': q.question_text,
            'type': q.question_type,
            'answer': q.correct_answer,
            'options': json.loads(q.options) if q.options else []
        }
        questions_data.append(q_data)
    
    # Path to template
    template_path = os.path.join(os.path.dirname(__file__), 'templates_files', 'exam_template.docx')
    
    # Check if template exists
    if not os.path.exists(template_path):
        # Fallback: create simple document without template
        flash('Template not found, creating simple document', 'warning')
        doc = Document()
        
        # Title
        title = doc.add_heading(exam.title, level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Info
        doc.add_paragraph(f'Number of Questions: {exam.num_questions}')
        doc.add_paragraph(f'Difficulty: {exam.difficulty.title()}')
        doc.add_paragraph(f'Date: {exam.created_at.strftime("%Y-%m-%d")}')
        doc.add_paragraph('_' * 60)
        
        # Questions
        for idx, q in enumerate(questions, 1):
            q_para = doc.add_paragraph()
            q_para.add_run(f'Question {idx}: ').bold = True
            q_para.add_run(q.question_text)
            
            if q.question_type == 'mcq' and q.options:
                options = json.loads(q.options)
                for i, option in enumerate(options, 1):
                    doc.add_paragraph(f'   {chr(64+i)}. {option}')
            
            ans_para = doc.add_paragraph()
            ans_para.add_run('Answer: ').bold = True
            ans_para.add_run(q.correct_answer)
            doc.add_paragraph()
        
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
    else:
        # Use template
        # Check how many versions are needed
        num_versions = getattr(exam, 'num_versions', 1) or 1
        teacher_name = getattr(exam, 'teacher_name', '') or ''
        show_answers = getattr(exam, 'show_answers', False)
        
        if num_versions > 1:
            zip_buffer = BytesIO()
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                for i in range(num_versions):
                    version_doc = DocxTemplate(template_path)
                    
                    # Shuffle questions and options for this version
                    import copy
                    shuffled_questions = copy.deepcopy(questions_data)
                    random.shuffle(shuffled_questions)
                    
                    # Shuffle options within MCQ questions
                    for sq in shuffled_questions:
                        if sq.get('type') == 'mcq' and sq.get('options'):
                            random.shuffle(sq['options'])
                    
                    context = {
                        'exam_title': f"{exam.title} - Model {chr(65+i)}", # Model A, B...
                        'exam_date': exam.created_at.strftime('%Y-%m-%d'),
                        'num_questions': exam.num_questions,
                        'difficulty': exam.difficulty.title(),
                        'questions': shuffled_questions,
                        'show_answers': show_answers,
                        'teacher_name': teacher_name,
                        'teacher': teacher_name # Typo fallback
                    }
                    version_doc.render(context)
                    
                    doc_buffer = BytesIO()
                    version_doc.save(doc_buffer)
                    zip_file.writestr(f'{exam.title}_Model_{chr(65+i)}.docx', doc_buffer.getvalue())
            
            zip_buffer.seek(0)
            return send_file(
                zip_buffer,
                as_attachment=True,
                download_name=f'{exam.title}_Models.zip',
                mimetype='application/zip'
            )
        else:
            doc = DocxTemplate(template_path)
            
            # Prepare context
            context = {
                'exam_title': exam.title,
                'exam_date': exam.created_at.strftime('%Y-%m-%d'),
                'num_questions': exam.num_questions,
                'difficulty': exam.difficulty.title(),
                'questions': questions_data,
                'show_answers': show_answers,
                'teacher_name': teacher_name,
                'teacher': teacher_name # Typo fallback
            }
            
            # Render template
            doc.render(context)
            
            # Save to buffer
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            return send_file(
                buffer,
                as_attachment=True,
                download_name=f'{exam.title}.docx',
                mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )


@app.route('/exam/<int:exam_id>/export/pdf')
@login_required
def export_pdf(exam_id):
    """Export exam to PDF format"""
    exam = Exam.query.get_or_404(exam_id)
    
    if exam.user_id != current_user.id:
        flash('Not authorized to export this exam', 'error')
        return redirect(url_for('dashboard'))
    
    questions = Question.query.filter_by(exam_id=exam_id).all()
    
    # Create PDF
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, 
        pagesize=letter,
        rightMargin=72,
        leftMargin=72,
        topMargin=72,
        bottomMargin=36
    )
    
    # Container for elements
    elements = []
    
    # Styles
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = styles['Heading1']
    title_style.fontSize = 16
    title_style.alignment = TA_LEFT
    title_style.spaceAfter = 20
    
    heading_style = styles['Heading2']
    heading_style.fontSize = 14
    heading_style.alignment = TA_LEFT
    heading_style.spaceAfter = 12
    
    normal_style = styles['Normal']
    normal_style.fontSize = 12
    normal_style.alignment = TA_LEFT
    normal_style.spaceAfter = 6
    
    # Title
    elements.append(Paragraph(exam.title, title_style))
    elements.append(Spacer(1, 0.2 * inch))
    
    # Exam info
    elements.append(Paragraph('<b>Exam Information</b>', heading_style))
    info_text = f'Number of Questions: {exam.num_questions}<br/>'
    info_text += f'Difficulty: {exam.difficulty.title()}<br/>'
    info_text += f'Date: {exam.created_at.strftime("%Y-%m-%d")}'
    elements.append(Paragraph(info_text, normal_style))
    elements.append(Spacer(1, 0.3 * inch))
    
    # Horizontal line
    elements.append(Paragraph('_' * 80, normal_style))
    elements.append(Spacer(1, 0.2 * inch))
    
    # Questions
    for idx, q in enumerate(questions, 1):
        # Question
        q_text = f'<b>Question {idx}:</b> {q.question_text}'
        elements.append(Paragraph(q_text, heading_style))
        elements.append(Spacer(1, 0.1 * inch))
        
        # Options for MCQ
        if q.question_type == 'mcq' and q.options:
            options = json.loads(q.options)
            for i, option in enumerate(options, 1):
                opt_text = f'   {chr(64+i)}. {option}'
                elements.append(Paragraph(opt_text, normal_style))
        
        # Answer
        ans_text = f'<b>Answer:</b> {q.correct_answer}'
        elements.append(Paragraph(ans_text, normal_style))
        elements.append(Spacer(1, 0.2 * inch))
    
    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'{exam.title}.pdf',
        mimetype='application/pdf'
    )


# إنشاء الجداول عند التشغيل لأول مرة
with app.app_context():
    db.create_all()


if __name__ == '__main__':
    app.run(debug=True)